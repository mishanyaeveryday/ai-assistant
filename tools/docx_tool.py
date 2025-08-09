# Creating Word files (python-docx)
import os
import json
from typing import List, Optional
from pydantic import BaseModel, Field
from docx import Document
from docx.shared import Inches
from langchain.tools import tool


class CreateDocxInput(BaseModel):
    title: str = Field(..., description="The main title of the document.")
    paragraphs: List[str] = Field(
        ..., description="A list of strings, where each string is a separate paragraph for the document.")
    save_path: str = Field(..., description="The full path, including the filename and .docx extension, where the document should be saved. For example: 'reports/my_report.docx'")
    table_data: Optional[List[List[str]]] = Field(
        None, description="Optional. Data for a table. It should be a list of lists, where the first inner list is the header row.")
    image_path: Optional[str] = Field(
        None, description="Optional. The local file path to an image to be included in the document.")


@tool(args_schema=CreateDocxInput)
def create_docx(
    title: str,
    paragraphs: List[str],
    save_path: str,
    table_data: Optional[List[List[str]]] = None,
    image_path: Optional[str] = None
) -> str:
    """
    Creates a new Word document (.docx) with a title, paragraphs, and optionally a table and an image. 
    Use this tool to save structured text, reports, or summaries into a file.
    """
    # Создание документа
    document = Document()

    if title:
        document.add_heading(title, level=0)

    if paragraphs:
        for para in paragraphs:
            document.add_paragraph(para)

    if table_data:
        if not table_data or not all(isinstance(row, list) for row in table_data):
            return "Error: table_data must be a list of lists."

        rows = len(table_data)
        cols = len(table_data[0]) if rows > 0 else 0
        if cols == 0:
            return "Error: table_data contains no columns."

        table = document.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'
        for i, row_data in enumerate(table_data):
            row_cells = table.rows[i].cells
            for j, cell_text in enumerate(row_data):
                row_cells[j].text = str(cell_text)

    if image_path:
        if os.path.exists(image_path):
            document.add_picture(image_path, width=Inches(4))
        else:
            print(f"Warning: Image path not found: {image_path}")

    output_dir = os.path.dirname(save_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    try:
        document.save(save_path)
        print(f"Document successfully created at: {save_path}")
        return f'Document successfully created at: "{save_path}"'
    except Exception as e:
        return f"Error saving document: {e}"
